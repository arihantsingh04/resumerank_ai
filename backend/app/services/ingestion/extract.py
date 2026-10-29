"""Text extraction from PDF / DOCX / TXT (D-07)."""

from __future__ import annotations

import io
import logging
import re

from app.core.errors import ValidationFailed

log = logging.getLogger(__name__)
MIN_TEXT_CHARS = 50
MAX_PDF_PAGES = 30


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValidationFailed("Password-protected PDFs aren't supported.")
    return "\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES])


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, kind: str) -> str:
    try:
        if kind == "pdf":
            raw = _pdf(data)
        elif kind == "docx":
            raw = _docx(data)
        else:
            raw = data.decode("utf-8", errors="replace")
    except ValidationFailed:
        raise
    except Exception as e:
        log.info("text extraction failed for %s file: %s", kind, type(e).__name__)
        raise ValidationFailed("Couldn't read this file. It may be corrupted.") from e
    text = _clean(raw)
    if len(text) < MIN_TEXT_CHARS:
        raise ValidationFailed("No extractable text. Scanned or image-only files aren't supported yet.")
    return text
